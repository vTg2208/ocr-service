"""Secure CSV/XLSX intake for legacy FRA registers.

The original workbook is retained as a private document. Each non-empty data
row becomes a reviewable archive record with field-level source evidence.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO, StringIO
import hashlib
from pathlib import Path
import re
import zipfile

from openpyxl import load_workbook
from sqlalchemy import select

from app.db.fra_completion_models import FRAImportBatch
from app.db.models import Document
from app.services.fra_archive import (
    ArchiveValidationError,
    create_archive_record,
    create_import_batch,
    process_archive_extraction,
)
from app.services.model_gateway import EntityExtractionResult
from app.services.state_profiles import get_state_profile


MAX_TABULAR_BYTES = 10 * 1024 * 1024
MAX_TABULAR_ROWS = 5000
MAX_TABULAR_COLUMNS = 100
MAX_CELL_CHARACTERS = 10_000
MAX_XLSX_ENTRIES = 2000
MAX_XLSX_EXPANDED_BYTES = 100 * 1024 * 1024


class TabularValidationError(ArchiveValidationError):
    pass


@dataclass(frozen=True)
class TabularUpload:
    filename: str
    content_type: str | None
    content: bytes


@dataclass(frozen=True)
class TabularRow:
    source_row: int
    source_sheet: str | None
    fields: dict
    evidence: dict


@dataclass(frozen=True)
class ParsedTabularUpload:
    safe_filename: str
    extension: str
    rows: list[TabularRow]


def _header_key(value: object) -> str:
    text = " ".join(str(value or "").strip().casefold().split())
    return re.sub(r"[^\w]+", "_", text, flags=re.UNICODE).strip("_")


HEADER_ALIASES = {
    "holder_name": {
        "holder_name", "claimant", "claimant_name", "applicant", "applicant_name",
        "patta_holder", "patta_holder_name", "rights_holder", "rights_holder_name",
    },
    "holder_type": {"holder_type", "claimant_type", "rights_holder_type"},
    "household_name": {"household", "household_name", "family_name"},
    "community_name": {"community", "community_name", "gram_sabha_name"},
    "state": {"state", "state_name"},
    "district": {"district", "district_name"},
    "block": {"block", "block_name", "taluk", "taluka", "tehsil"},
    "village": {"village", "village_name", "revenue_village"},
    "survey_number": {"survey_number", "survey_no", "survey_num", "sy_no"},
    "subdivision_number": {
        "subdivision_number", "sub_division_number", "subdivision_no", "sub_division_no",
    },
    "right_type": {"right_type", "claim_type", "fra_right", "fra_type"},
    "claim_status": {"claim_status", "status", "application_status"},
    "claim_number": {"claim_number", "claim_no", "application_number", "application_no"},
    "claim_year": {"claim_year", "application_year", "year"},
    "claimed_area": {"claimed_area", "area_claimed"},
    "claimed_area_sqm": {
        "claimed_area_sqm", "claimed_area_sq_m", "claimed_area_m2", "area_claimed_sqm",
    },
    "granted_area": {"granted_area", "area_granted"},
    "granted_area_sqm": {
        "granted_area_sqm", "granted_area_sq_m", "granted_area_m2", "area_granted_sqm",
    },
    "area_unit": {"area_unit", "unit", "land_area_unit"},
    "gram_sabha_status": {"gram_sabha", "gram_sabha_status", "gs_status"},
    "sdlc_status": {"sdlc", "sdlc_status"},
    "dlc_status": {"dlc", "dlc_status"},
    "decision_date": {"decision_date", "order_date", "approval_date"},
    "title_number": {"title_number", "title_no", "title_reference", "patta_number", "patta_no"},
    "latitude": {"latitude", "lat"},
    "longitude": {"longitude", "lon", "lng", "long"},
    "coordinates": {"coordinates", "coordinate", "lat_long", "lat_lon"},
}
ALIAS_TO_FIELD = {
    alias: field_name for field_name, aliases in HEADER_ALIASES.items() for alias in aliases
}


def _safe_filename(filename: str) -> tuple[str, str]:
    basename = Path(str(filename or "").replace("\\", "/")).name
    basename = re.sub(r"[^\w.() -]+", "_", basename, flags=re.UNICODE).strip(" .")
    if not basename:
        raise TabularValidationError("A CSV or XLSX filename is required.")
    extension = Path(basename).suffix.casefold().lstrip(".")
    if extension not in {"csv", "xlsx"}:
        raise TabularValidationError("Legacy tabular intake accepts CSV or XLSX files only.")
    return basename[:255], extension


def _scalar(value: object) -> object:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, str):
        cleaned = " ".join(value.split())
        if len(cleaned) > MAX_CELL_CHARACTERS:
            raise TabularValidationError("A tabular cell exceeds the maximum supported length.")
        return cleaned or None
    if isinstance(value, (bool, int, float)):
        return value
    cleaned = " ".join(str(value).split())
    if len(cleaned) > MAX_CELL_CHARACTERS:
        raise TabularValidationError("A tabular cell exceeds the maximum supported length.")
    return cleaned or None


def _mapped_headers(headers: list[object]) -> list[tuple[str, str]]:
    if not headers or not any(_header_key(value) for value in headers):
        raise TabularValidationError("The tabular source must contain a header row.")
    if len(headers) > MAX_TABULAR_COLUMNS:
        raise TabularValidationError("The tabular source has too many columns.")
    mapped = []
    seen_keys = set()
    seen_fields = set()
    for position, source_header in enumerate(headers, start=1):
        key = _header_key(source_header)
        if not key:
            raise TabularValidationError(f"Header column {position} is empty.")
        if key in seen_keys:
            raise TabularValidationError(f"Duplicate header: {source_header}.")
        field_name = ALIAS_TO_FIELD.get(key, f"legacy_{key}")
        if field_name in seen_fields:
            raise TabularValidationError(
                f"Multiple headers map to the same FRA field: {field_name}."
            )
        seen_keys.add(key)
        seen_fields.add(field_name)
        mapped.append((field_name[:100], str(source_header).strip()))
    return mapped


def _build_rows(raw_rows, *, source_sheet: str | None = None) -> list[TabularRow]:
    iterator = iter(raw_rows)
    try:
        headers = list(next(iterator))
    except StopIteration as error:
        raise TabularValidationError("The tabular source is empty.") from error
    mapped = _mapped_headers(headers)
    rows = []
    for source_row, raw_values in enumerate(iterator, start=2):
        values = list(raw_values)
        if len(values) > len(mapped) and any(_scalar(value) is not None for value in values[len(mapped):]):
            raise TabularValidationError(f"Row {source_row} contains values beyond the header columns.")
        values.extend([None] * (len(mapped) - len(values)))
        if not any(_scalar(value) is not None for value in values[:len(mapped)]):
            continue
        if len(rows) >= MAX_TABULAR_ROWS:
            raise TabularValidationError(
                f"A tabular source can contain at most {MAX_TABULAR_ROWS} data rows."
            )
        fields = {}
        evidence = {}
        for (field_name, source_header), raw_value in zip(mapped, values):
            value = _scalar(raw_value)
            if value is None:
                continue
            fields[field_name] = value
            evidence[field_name] = {
                "source_row": source_row,
                "source_sheet": source_sheet,
                "source_header": source_header,
                "source_value": value,
                "extraction_method": "tabular_header_mapping",
                "confidence": 1.0,
            }
        rows.append(TabularRow(source_row, source_sheet, fields, evidence))
    if not rows:
        raise TabularValidationError("The tabular source contains no data rows.")
    return rows


def _csv_rows(content: bytes):
    if b"\x00" in content:
        raise TabularValidationError("The CSV file contains binary data.")
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise TabularValidationError("CSV files must use UTF-8 encoding.") from error
    try:
        return _build_rows(csv.reader(StringIO(text, newline="")))
    except csv.Error as error:
        raise TabularValidationError("The CSV file is malformed.") from error


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            names = {entry.filename.replace("\\", "/") for entry in entries}
            if not 1 <= len(entries) <= MAX_XLSX_ENTRIES:
                raise TabularValidationError("The XLSX archive entry count is invalid.")
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise TabularValidationError("The upload is not a valid XLSX workbook.")
            if any(
                name.startswith("/") or ".." in Path(name).parts
                for name in names
            ):
                raise TabularValidationError("The XLSX archive contains an unsafe path.")
            if sum(entry.file_size for entry in entries) > MAX_XLSX_EXPANDED_BYTES:
                raise TabularValidationError("The XLSX workbook expands beyond 100 MB.")
    except zipfile.BadZipFile as error:
        raise TabularValidationError("The upload is not a valid XLSX workbook.") from error


def _xlsx_rows(content: bytes):
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as error:
        raise TabularValidationError("The XLSX workbook could not be read.") from error
    try:
        sheet = workbook.active
        raw_rows = []
        for row_number, cells in enumerate(sheet.iter_rows(), start=1):
            if any(cell.data_type == "f" for cell in cells):
                raise TabularValidationError(
                    f"The XLSX workbook contains a formula on row {row_number}; provide values only."
                )
            raw_rows.append([cell.value for cell in cells])
        return _build_rows(raw_rows, source_sheet=sheet.title)
    finally:
        workbook.close()


def parse_tabular_upload(upload: TabularUpload) -> ParsedTabularUpload:
    safe_filename, extension = _safe_filename(upload.filename)
    if not upload.content:
        raise TabularValidationError("The tabular source is empty.")
    if len(upload.content) > MAX_TABULAR_BYTES:
        raise TabularValidationError("The tabular source exceeds 10 MB.")
    rows = _csv_rows(upload.content) if extension == "csv" else _xlsx_rows(upload.content)
    return ParsedTabularUpload(safe_filename, extension, rows)


class TabularFRAExtractor:
    version = "fra-tabular-schema-v1"

    def extract(self, document_reference: str, manifest: dict) -> EntityExtractionResult:
        fields = dict(manifest["fields"])
        evidence = dict(manifest["evidence"])
        profile = get_state_profile(fields.get("state") or "TN")
        if "state" not in fields:
            fields["state"] = profile.name
            evidence["state"] = {
                "source_row": manifest["source_row"],
                "source_sheet": manifest.get("source_sheet"),
                "source_value": None,
                "extraction_method": "intake_context",
                "confidence": 1.0,
            }
        fields["state_code"] = profile.code
        evidence["state_code"] = {
            "source_row": manifest["source_row"],
            "source_sheet": manifest.get("source_sheet"),
            "source_value": fields.get("state"),
            "extraction_method": "state_profile_mapping",
            "confidence": 1.0,
        }
        for field_name, normalizer in (
            ("district", profile.normalize_district),
            ("block", profile.normalize_block),
            ("village", profile.normalize_village),
        ):
            if fields.get(field_name) is not None:
                fields[field_name] = normalizer(str(fields[field_name]))
        if not fields.get("holder_name"):
            for source_name in ("household_name", "community_name"):
                if fields.get(source_name):
                    fields["holder_name"] = fields[source_name]
                    evidence["holder_name"] = {
                        **dict(evidence.get(source_name) or {}),
                        "extraction_method": f"{source_name}_alias",
                    }
                    break
        if fields.get("right_type") is not None:
            fields["right_type"] = str(fields["right_type"]).strip().upper()
        if fields.get("claim_status") is not None:
            fields["claim_status"] = str(fields["claim_status"]).strip().casefold()
        return EntityExtractionResult(
            fields=fields,
            field_evidence=evidence,
            confidence=None,
            model_version=self.version,
            processing_time_ms=0,
            provenance={
                "adapter": "tabular_header_mapping",
                "synthetic": False,
                "document_reference": document_reference,
                "source_row": manifest["source_row"],
                "source_sheet": manifest.get("source_sheet"),
            },
        )


def _existing_result(batch: FRAImportBatch) -> dict:
    records = [
        {
            "record_id": str(record.id),
            "legacy_reference": record.legacy_reference,
            "source_row": (record.provenance_json or {}).get("source_row"),
            "status": "accepted",
        }
        for record in sorted(batch.records, key=lambda item: ((item.provenance_json or {}).get("source_row", 0), str(item.id)))
    ]
    return {
        "batch_id": str(batch.id),
        "batch_status": batch.status,
        "accepted": len(records),
        "rejected": batch.failed_count,
        "replayed": True,
        "records": records,
    }


def ingest_tabular_archive(
    session,
    *,
    upload: TabularUpload,
    source_office: str,
    district: str,
    actor_id,
    idempotency_key: str,
    storage,
    scanner,
    state: str = "Tamil Nadu",
    request_id: str | None = None,
) -> dict:
    source = " ".join(str(source_office or "").split())
    district_name = " ".join(str(district or "").split())
    key = " ".join(str(idempotency_key or "").split())
    if not source or not district_name or not key:
        raise TabularValidationError("Source office, district, and Idempotency-Key are required.")
    existing = session.scalar(
        select(FRAImportBatch).where(
            FRAImportBatch.created_by == actor_id,
            FRAImportBatch.idempotency_key == key,
        )
    )
    if existing is not None:
        return _existing_result(existing)

    scanner.scan(upload.content)
    parsed = parse_tabular_upload(upload)
    checksum = hashlib.sha256(upload.content).hexdigest()
    duplicate = session.scalar(
        select(Document.id).where(
            Document.uploaded_by == actor_id,
            Document.sha256 == checksum,
        ).limit(1)
    )
    if duplicate:
        raise TabularValidationError("This source has already been submitted by the current uploader.")

    storage_key = storage.put(upload.content, f".{parsed.extension}")
    try:
        with session.begin_nested():
            batch = create_import_batch(
                session,
                source_label=source,
                state=state,
                actor_id=actor_id,
                idempotency_key=key,
                synthetic=False,
                provenance={
                    "source": source,
                    "source_office": source,
                    "district": district_name,
                    "synthetic": False,
                    "ingest_method": "tabular_upload",
                    "filename": parsed.safe_filename,
                    "sha256": checksum,
                },
                request_id=request_id,
            )
            document = Document(
                uploaded_by=actor_id,
                storage_key=storage_key,
                original_filename=parsed.safe_filename,
                content_type=(
                    "text/csv" if parsed.extension == "csv"
                    else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                ),
                sha256=checksum,
                ocr_status="completed",
                idempotency_key=f"fra-tabular:{batch.id}:{checksum[:16]}",
            )
            session.add(document)
            session.flush()
            stem = Path(parsed.safe_filename).stem.strip("._-") or "legacy-register"
            results = []
            extractor = TabularFRAExtractor()
            for row in parsed.rows:
                reference = f"{stem}:row:{row.source_row}"
                record = create_archive_record(
                    session,
                    batch=batch,
                    document_id=document.id,
                    legacy_reference=reference,
                    actor_id=actor_id,
                    synthetic=False,
                    provenance={
                        **dict(batch.provenance_json),
                        "source_row": row.source_row,
                        "source_sheet": row.source_sheet,
                    },
                    request_id=request_id,
                )
                process_archive_extraction(
                    session,
                    record,
                    extractor=extractor,
                    manifest={
                        "fields": row.fields,
                        "evidence": row.evidence,
                        "source_row": row.source_row,
                        "source_sheet": row.source_sheet,
                    },
                    actor_id=actor_id,
                    request_id=request_id,
                )
                results.append(
                    {
                        "record_id": str(record.id),
                        "legacy_reference": reference,
                        "source_row": row.source_row,
                        "status": "accepted",
                    }
                )
            session.flush()
    except Exception:
        storage.delete(storage_key)
        raise

    return {
        "batch_id": str(batch.id),
        "batch_status": batch.status,
        "accepted": len(results),
        "rejected": 0,
        "replayed": False,
        "records": results,
    }


__all__ = [
    "ParsedTabularUpload",
    "TabularRow",
    "TabularUpload",
    "TabularValidationError",
    "ingest_tabular_archive",
    "parse_tabular_upload",
]
